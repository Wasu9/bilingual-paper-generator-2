"""PaperCraft AI Phase 5: transparent Word layout hardening.

Python imports sitecustomize automatically when the repository root is on sys.path.
This module deliberately post-processes only python-docx documents created by the app;
it does not alter the PaperCraft UI or parsing/translation/visual pipelines.
"""
from docx import document as _document_module
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_ORIGINAL_SAVE = _document_module.Document.save


def _set_cell_margins(cell, top=55, start=65, bottom=55, end=65):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:tblHeader')) is None:
        trPr.append(OxmlElement('w:tblHeader'))


def _fixed_table(table):
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    table.autofit = False


def _paragraph_quality(paragraph, keep_with_next=False):
    pf = paragraph.paragraph_format
    pf.widow_control = True
    pf.keep_with_next = keep_with_next
    if pf.space_after is None:
        pf.space_after = Pt(2)
    for r in paragraph.runs:
        if r.font.size is None:
            r.font.size = Pt(9.5)


def _harden_document(doc):
    for section in doc.sections:
        # Keep the existing compact paper margins, but make header/footer clearance explicit.
        section.header_distance = Pt(10)
        section.footer_distance = Pt(10)

    for table in doc.tables:
        _fixed_table(table)
        if table.rows:
            _repeat_header(table.rows[0])

        # A question is one logical unit: prevent Word from cutting the table row itself.
        # Word can still move an oversized row to the next page rather than corrupting it.
        for row in table.rows:
            _cant_split(row)
            for cell in row.cells:
                _set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    # Keep the question line attached to following option paragraphs.
                    text = (paragraph.text or '').strip()
                    is_question = bool(text[:4].split('.')[0].isdigit() and '.' in text[:5])
                    _paragraph_quality(paragraph, keep_with_next=is_question)
                    # Prevent accidental blank paragraphs from consuming layout space.
                    if not text:
                        paragraph.paragraph_format.space_after = Pt(0)

    # Apply a conservative default to all runs without overriding the app's Hindi font.
    for paragraph in doc.paragraphs:
        _paragraph_quality(paragraph)


def _patched_save(self, path_or_stream):
    _harden_document(self)
    return _ORIGINAL_SAVE(self, path_or_stream)


# Patch the class method once. The app's existing `from docx import Document`
# resolves to the same python-docx Document implementation.
if _document_module.Document.save is not _patched_save:
    _document_module.Document.save = _patched_save
