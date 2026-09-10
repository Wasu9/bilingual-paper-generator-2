import io,re,base64,os,requests
import fitz,streamlit as st
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from deep_translator import GoogleTranslator

st.set_page_config(page_title="PaperCraft AI — Bilingual Paper Generator",page_icon="🎓",layout="wide",initial_sidebar_state="collapsed")

st.markdown(r"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
.stApp{background:#f8fafc;font-family:Inter,sans-serif;color:#0f172a}.block-container{max-width:1180px;padding:1.2rem 2rem 4rem}header[data-testid="stHeader"]{background:transparent}.nav{display:flex;align-items:center;justify-content:space-between;padding:10px 0 22px}.brand{display:flex;align-items:center;gap:10px;font-weight:800;font-size:20px;color:#0f172a}.brandmark{width:38px;height:38px;border-radius:11px;background:linear-gradient(135deg,#2563eb,#7c3aed);display:grid;place-items:center;color:white;font-size:19px;box-shadow:0 8px 22px #2563eb33}.navlinks{display:flex;gap:28px;color:#64748b;font-size:14px;font-weight:600}
.hero-wrap{border-radius:30px;padding:68px 58px 60px;background:radial-gradient(circle at 85% 15%,#60a5fa55 0,transparent 28%),radial-gradient(circle at 10% 90%,#a78bfa44 0,transparent 30%),linear-gradient(135deg,#0f172a,#1e3a8a 58%,#312e81);color:#fff;box-shadow:0 24px 60px #0f172a20;overflow:hidden;position:relative}.hero-wrap:after{content:"";position:absolute;width:300px;height:300px;border:1px solid #ffffff18;border-radius:50%;right:-110px;bottom:-140px}.eyebrow{display:inline-block;background:#ffffff14;border:1px solid #ffffff22;border-radius:999px;padding:8px 14px;font-size:12px;font-weight:700;letter-spacing:.3px;margin-bottom:20px}.hero-title{font-size:52px;line-height:1.05;letter-spacing:-2px;font-weight:800;max-width:780px;margin:0 0 18px}.hero-title span{color:#93c5fd}.hero-sub{font-size:17px;line-height:1.7;color:#dbeafe;max-width:700px;margin-bottom:28px}.hero-points{display:flex;gap:18px;flex-wrap:wrap;font-size:13px;color:#e0e7ff}.hero-points span{background:#ffffff0d;border:1px solid #ffffff18;padding:9px 12px;border-radius:10px}
.section{padding:42px 0 8px;text-align:center}.kicker{color:#2563eb;font-size:12px;font-weight:800;letter-spacing:1.3px;text-transform:uppercase}.section h2{font-size:30px;letter-spacing:-.8px;margin:8px 0 10px;color:#0f172a}.section p{color:#64748b;max-width:650px;margin:0 auto 28px;line-height:1.6}.feature{background:#fff;border:1px solid #e2e8f0;border-radius:18px;padding:24px;text-align:left;height:100%;box-shadow:0 8px 28px #0f172a08}.feature-icon{width:42px;height:42px;border-radius:12px;background:#eff6ff;display:grid;place-items:center;font-size:20px;margin-bottom:15px}.feature h3{margin:0 0 8px;font-size:16px}.feature p{margin:0;color:#64748b;font-size:13px;line-height:1.55}.step{background:#fff;border:1px solid #e2e8f0;border-radius:18px;padding:22px;text-align:center;height:100%}.stepno{font-size:12px;font-weight:800;color:#2563eb;background:#eff6ff;border-radius:999px;padding:6px 10px;display:inline-block;margin-bottom:10px}.step h3{font-size:16px;margin:5px}.step p{font-size:13px;color:#64748b;margin:0;line-height:1.5}
.upload-shell{background:#fff;border:1px solid #dbe4f0;border-radius:24px;padding:30px;box-shadow:0 16px 45px #0f172a0b;margin-top:20px}.upload-title{text-align:center;font-size:24px;font-weight:800;margin:0 0 7px}.upload-sub{text-align:center;color:#64748b;font-size:14px;margin-bottom:22px}.status-pill{display:inline-block;border-radius:999px;padding:7px 12px;background:#ecfdf5;color:#047857;font-size:12px;font-weight:700}.footer{text-align:center;color:#94a3b8;font-size:12px;padding:42px 0 8px}.stButton>button{border-radius:12px!important;height:46px!important;font-weight:700!important;background:linear-gradient(135deg,#2563eb,#4f46e5)!important;border:0!important;box-shadow:0 8px 20px #4f46e533!important}.stDownloadButton>button{border-radius:12px!important;height:46px!important;font-weight:700!important;background:linear-gradient(135deg,#059669,#047857)!important;border:0!important;box-shadow:0 8px 20px #05966933!important;color:white!important}[data-testid="stFileUploader"]{background:#f8fafc;border:1.5px dashed #93c5fd;border-radius:18px;padding:18px}[data-testid="stFileUploaderDropzone"]{background:#fff;border-radius:14px}#MainMenu,footer{visibility:hidden}@media(max-width:800px){.hero-wrap{padding:42px 25px}.hero-title{font-size:37px}.navlinks{display:none}.block-container{padding:1rem}}
</style>
""",unsafe_allow_html=True)

st.markdown("""<div class="nav"><div class="brand"><div class="brandmark">✦</div>PaperCraft AI</div><div class="navlinks"><span>Features</span><span>How it works</span><span>Bilingual Papers</span></div><div><span class="status-pill">● AI Paper Generator</span></div></div>""",unsafe_allow_html=True)
st.markdown("""<div class="hero-wrap"><div class="eyebrow">SMARTER PAPER CREATION • ENGLISH → ENGLISH + HINDI</div><h1 class="hero-title">Turn any question paper into a <span>professional bilingual paper.</span></h1><div class="hero-sub">Upload your English PDF and PaperCraft AI extracts questions, protects scientific notation, translates the content into Hindi, and prepares a clean two-column Word document.</div><div class="hero-points"><span>✓ Question extraction</span><span>✓ Hindi translation</span><span>✓ Math & science protection</span><span>✓ Word-ready formatting</span></div></div>""",unsafe_allow_html=True)
st.markdown('<div class="section"><div class="kicker">Built for educators</div><h2>Everything you need to create bilingual papers</h2><p>A clean workflow designed for coaching institutes, teachers and academic teams who need consistent English + Hindi question papers.</p></div>',unsafe_allow_html=True)
f1,f2,f3=st.columns(3)
with f1: st.markdown('<div class="feature"><div class="feature-icon">📄</div><h3>Smart PDF Extraction</h3><p>Detects numbered questions and cleans common headers, page numbers and repeated paper metadata.</p></div>',unsafe_allow_html=True)
with f2: st.markdown('<div class="feature"><div class="feature-icon">🧠</div><h3>Science-Aware Translation</h3><p>Protects formulas, units, symbols and common scientific terms before Hindi translation.</p></div>',unsafe_allow_html=True)
with f3: st.markdown('<div class="feature"><div class="feature-icon">📝</div><h3>Professional Word Output</h3><p>Generates a polished two-column English + Hindi layout ready for editing and printing.</p></div>',unsafe_allow_html=True)
st.markdown('<div class="section"><div class="kicker">Simple workflow</div><h2>From PDF to paper in three steps</h2></div>',unsafe_allow_html=True)
s1,s2,s3=st.columns(3)
with s1: st.markdown('<div class="step"><div class="stepno">01</div><h3>Upload</h3><p>Choose your English question-paper PDF.</p></div>',unsafe_allow_html=True)
with s2: st.markdown('<div class="step"><div class="stepno">02</div><h3>Generate</h3><p>Questions are extracted, translated and formatted.</p></div>',unsafe_allow_html=True)
with s3: st.markdown('<div class="step"><div class="stepno">03</div><h3>Download</h3><p>Get the finished bilingual Word document.</p></div>',unsafe_allow_html=True)

HEAD=('PHYSICS','CHEMISTRY','MATHEMATICS','BIOLOGY','BOTANY','ZOOLOGY','I PUC','II PUC','JEE MAINS','NEET','QUESTION PAPER','TEST SERIES','ANSWER KEY','SYLLABUS COVERED','IMPORTANT INSTRUCTIONS FOR CANDIDATES','STUDENT INFORMATION')
PAGE=re.compile(r'^(?:VERSION\s*[- ]?\s*\d+\s+)?I\s*[- ]*PUC.*?Page\s*\d+\s+of\s+\d+$',re.I)
QUESTION_START=re.compile(r'^\s*(\d{1,3})(?:\.(?!\d)\s*|\s+(?=(?:What|Which|When|Where|Who|Why|How|A|An|The|If|During|Identify|Choose|Select|Calculate|Find|Determine|Consider|Given|For|In|On|According|Assertion|Reason|Match|Name|Our|Cells|Cell|Gas|Resistance|Eukaryotic|Mechanisms|Movement|Additional|Tidal|Normal|Inspiratory|Total|Prophase|Chromosomes|Meiosis|Does|Is|Are|Can|Let)\b))',re.I|re.M)
PAGE_INLINE=re.compile(r'(?i)\bPage\s*\d+\s+of\s+\d+\b|\bVERSION\s*[- ]?\s*\d+\s*[/\-]?\s*I\s*[- ]*PUC\s*(?:[- ]*(?:NEET|JEE\s*MAINS))?\s*(?:[/\-]?\s*of\s*\d+)?')
SEC=re.compile(r'^\(?\s*(single correct|multiple correct|numerical value|assertion|paragraph|match|matrix)\b',re.I)
LIG={'Ư':'ff','ﬁ':'fi','ﬂ':'fl','ﬀ':'ff','ﬃ':'ffi','ﬄ':'ffl'}
PROTECT=re.compile('\\\\(?:[A-Za-z]+(?:\\{[^{}]*\\})?)|\\$[^$]+\\$|\\b(?:sin|cos|tan|cot|sec|cosec|log|ln|lim|exp|dx|dy|dt|dA|dB|dV|vec|frac)\\b|\\b(?:H2O|CO2|O2|N2|H2|NaCl|HCl|H2SO4|HNO3|NH3|CH4|C2H5OH|KMnO4|K2Cr2O7|NaOH|CaCO3|CaO|CH3MgBr|LiAlH4|NaBH4|ATP|DNA|mRNA|tRNA|eV|keV|MeV|GeV|Hz|kHz|MHz|GHz|Pa|kPa|MPa|mol|kg|mg|μg|cm|mm|nm|km|ms|μs)\\b|\\b[A-Za-zα-ωΑ-Ω]+(?:\\^[-+]?\\d+|[₀₁₂₃₄₅₆₇₈₉])+\\b|\\b[A-Za-zα-ωΑ-Ω0-9]+\\s*[=≈≠≤≥<>±∝→←↔⇌]\\s*[A-Za-zα-ωΑ-Ω0-9²³⁴⁵⁻⁺⁰¹²³⁴⁵⁶⁷⁸⁹./()+×÷√-]+(?:\\s+[A-Za-zα-ωΑ-Ω0-9²³⁴⁵⁻⁺⁰¹²³⁴⁵⁶⁷⁸⁹./()+×÷√-]+)*|[∫∮∑√∞±×÷≈≠≤≥→←↔⇌∆∂∇θπλμΩαβγδφψω²³⁴⁵⁻⁺⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉]',re.I)

def norm(x):
    for a,b in LIG.items(): x=x.replace(a,b)
    x=x.replace('−','-').replace('–','-').replace('—','-')
    return re.sub(r'\s+',' ',x.replace('\xa0',' ')).strip()

@st.cache_data(show_spinner=False)
def pages(data):
    p=fitz.open(stream=data,filetype='pdf');return [x.get_text('text') for x in p]

def clean(x): return norm(x)

def math(x):
    m={'𝑥':'x','𝑦':'y','𝑧':'z','𝑛':'n','𝑎':'a','𝑏':'b','𝑚':'m','𝑝':'p','𝑞':'q','𝑟':'r','𝑡':'t','𝑓':'f','𝑒':'e','𝑑':'d','𝑐':'c','𝛼':'α','𝛽':'β','𝛾':'γ','𝛿':'δ','𝜃':'θ','𝜋':'π','𝜆':'λ','𝜇':'μ','𝜔':'ω','𝛺':'Ω'}
    for a,b in m.items(): x=x.replace(a,b)
    x=x.replace('∗','×').replace('−','-').replace('–','-').replace('—','-')
    x=re.sub(r'\s+',' ',x).strip()
    x=re.sub(r'\s*([=+\-×÷≤≥≠<>∝→←↔⇌])\s*',r' \1 ',x)
    return re.sub(r'\s+',' ',x).strip()
def fmt(lines):
    a=[math(x) for x in lines if x.strip()];o=[];i=0
    while i<len(a):
        if i+2<len(a) and re.fullmatch(r'[\d.]+',a[i+1]) and re.fullmatch(r'[\d.]+',a[i+2]) and a[i].rstrip().endswith(','):
            o.append(f'{a[i]} {a[i+1]}/{a[i+2]}');i+=3;continue
        if i+1<len(a) and re.fullmatch(r'[\d.]+',a[i]) and re.fullmatch(r'[\d.]+(?:\s*[A-Za-z²³⁴⁵⁻−0-9]+)?',a[i+1]):
            o.append(f'{a[i]}/{a[i+1]}');i+=2;continue
        if a[i] in ('∫','∮') and i+2<len(a) and a[i+1] in ('dx','dy','dt','dA','dV') and a[i+2].endswith('='):
            o.append(f'{a[i]} {a[i+1]}/({a[i+2][:-1].strip()}) =');i+=3;continue
        if i+1<len(a):
            b=a[i+1];c=re.sub(r'\s+','',b);z=re.match(r'^([\d.,a-zA-Zα-ω]+)\s*\+\s*(c\.?)$',b)
            if z and len(z.group(1))<15:o.append(f'({a[i]})/({z.group(1)}) + {z.group(2)}');i+=2;continue
            if re.fullmatch(r'[\d.,a-zA-Zα-ω+\-−=]+',c) and len(c)<15 and re.search(r'[\dA-Za-zα-ω}]',a[i]) and not a[i].endswith(('=',':')):
                o.append(f'({a[i]})/({c})');i+=2;continue
        o.append(a[i]);i+=1
    s=' '.join(o).strip();s=re.sub(r'\(([^()]{1,80})\)/\(([^()]{1,40})\)',r'\1/\2',s)
    return re.sub(r'\s+',' ',s).strip()

def _question_candidates(text):
    return [(m.start(),int(m.group(1))) for m in QUESTION_START.finditer(text) if 1 <= int(m.group(1)) <= 500]

def blocks(ps):
    t='\n'.join(ps); candidates=_question_candidates(t); runs=[]
    for i,(pos,num) in enumerate(candidates):
        if num != 1: continue
        run=[(pos,num)]; last=num
        for pos2,num2 in candidates[i+1:]:
            if num2 == last+1:
                run.append((pos2,num2)); last=num2; continue
            if last >= 10 and 1 <= num2 <= 4: continue
            break
        if len(run) >= 3:
            first_end=run[1][0] if len(run)>1 else len(t)
            first_block=t[pos:first_end]
            option_count=len(re.findall(r'\((?:1|2|3|4)\)',first_block))
            question_words=len(re.findall(r'\b(?:what|which|when|where|who|why|how|find|calculate|given|consider|identify|determine|select|choose|the|during|assertion|reason)\b',first_block,re.I))
            runs.append(((len(run),option_count,question_words,len(first_block)),run))
    if not runs: return []
    best=max(runs,key=lambda z:z[0])[1]
    return [(num,t[pos:best[i+1][0] if i+1<len(best) else len(t)]) for i,(pos,num) in enumerate(best)]

def parse(n,b):
    ls=[]
    for x in b.splitlines():
        x=norm(x); x=PAGE_INLINE.sub(' ',x)
        x=re.sub(r'\s*/\s*$', '', x)
        x=re.sub(r'\s*\(?/?\s*correct\.\s*\)?\s*$', '', x, flags=re.I)
        x=norm(x)
        if not x or PAGE.fullmatch(x) or x.upper() in HEAD or SEC.match(x) or x.lower().startswith(('this section contains','marking scheme:')): continue
        ls.append(x)
    if ls:
        qprefix=re.compile(rf'^\s*{n}(?:\.(?!\d))?\s*',re.I)
        ls[0]=qprefix.sub('',ls[0],1).strip()
    parts={1:[],2:[],3:[],4:[]}; stem=[]; cur=None
    for x in ls:
        markers=list(re.finditer(r'\((1|2|3|4)\)',x))
        if markers:
            prefix=x[:markers[0].start()].strip()
            if prefix: (parts[cur] if cur is not None else stem).append(prefix)
            for j,m in enumerate(markers):
                k=int(m.group(1)); v=x[m.end():markers[j+1].start() if j+1<len(markers) else len(x)].strip()
                if v: parts[k].append(v)
                cur=k
        elif cur is None: stem.append(x)
        else: parts[cur].append(x)
    return {'num':n,'stem':fmt(stem),'options':[fmt(parts[k]) for k in (1,2,3,4)]}

def validate_records(rs):
    nums=[q['num'] for q in rs]; issues=[]
    if not nums: return ['No numbered questions detected.']
    if nums[0] != 1: issues.append(f'Question sequence starts at {nums[0]}, not 1.')
    missing=[n for n in range(1,max(nums)+1) if n not in nums]
    duplicates=sorted({n for n in nums if nums.count(n)>1})
    if missing: issues.append('Missing question number(s): '+', '.join(map(str,missing[:20])))
    if duplicates: issues.append('Duplicate question number(s): '+', '.join(map(str,duplicates[:20])))
    if nums != sorted(nums): issues.append('Question numbers are out of order.')
    return issues

def records(ps): return [parse(n,b) for n,b in blocks(ps)]

def _spans(page):
    out=[]
    for b in page.get_text('dict').get('blocks',[]):
        for line in b.get('lines',[]):
            for s in line.get('spans',[]):
                tx=s.get('text','').strip()
                if tx:out.append((s['bbox'],tx))
    return sorted(out,key=lambda z:(z[0][1],z[0][0]))

def _q_ranges(page):
    qs=[];seen=set()
    for bbox,tx in _spans(page):
        m=re.match(r'^(\d{1,3})\.\s*',tx)
        if m and int(m.group(1)) not in seen:seen.add(int(m.group(1)));qs.append((int(m.group(1)),bbox[1],bbox[3]))
    return qs

def _option_markers(page,qy0,next_y):
    d={}
    for bbox,tx in _spans(page):
        if not (qy0<=bbox[1]<next_y+2):continue
        for m in re.finditer(r'\((1|2|3|4)\)',tx):
            k=int(m.group(1));d.setdefault(k,(k,bbox[0],bbox[1],bbox[2],bbox[3]))
    return d

def _small_images(page):
    items=[];pw,ph=page.rect.width,page.rect.height
    for im in page.get_images(full=True):
        try:r=page.get_image_rects(im[0])[0]
        except Exception:continue
        if r.width*r.height>pw*ph*.16 or r.width<18 or r.height<12:continue
        items.append((r,im[0]))
    return items

def _pix_crop(page,rect,pad=2):
    r=fitz.Rect(max(page.rect.x0,rect.x0-pad),max(page.rect.y0,rect.y0-pad),min(page.rect.x1,rect.x1+pad),min(page.rect.y1,rect.y1+pad))
    return page.get_pixmap(matrix=fitz.Matrix(2,2),clip=r,alpha=False).tobytes('png')

def _marker_centers(markers):
    return {k:((v[1]+v[3])/2,(v[2]+v[4])/2) for k,v in markers.items()}

def _crop_image(page,rect,pad=2,scale=2):
    r=fitz.Rect(max(page.rect.x0,rect.x0-pad),max(page.rect.y0,rect.y0-pad),min(page.rect.x1,rect.x1+pad),min(page.rect.y1,rect.y1+pad))
    return page.get_pixmap(matrix=fitz.Matrix(scale,scale),clip=r,alpha=False).tobytes('png')

def _split_composite_option(page,rect,markers):
    centers=_marker_centers(markers)
    if len(centers)<3 or rect.width<page.rect.width*.45:return {}
    inside=[k for k,(x,y) in centers.items() if rect.x0-30<=x<=rect.x1+30 and rect.y0-3<=y<=rect.y1+3]
    if len(inside)<3:return {}
    keys=sorted(inside,key=lambda k:centers[k][0])
    if keys != sorted(centers):return {}
    xs=[centers[k][0] for k in keys]
    bounds=[rect.x0]+[(xs[i]+xs[i+1])/2 for i in range(len(xs)-1)]+[rect.x1]
    out={}
    for k,x0,x1 in zip(keys,bounds,bounds[1:]):
        if x1-x0<18:continue
        out[k]=_crop_image(page,fitz.Rect(x0,rect.y0,x1,rect.y1),pad=1)
    return out if len(out)>=3 else {}

def _assign_option_images(candidates,markers):
    centers=_marker_centers(markers);out={};used=set()
    for rect,img in sorted(candidates,key=lambda z:(z[0].y0,z[0].x0)):
        possible=[]
        for k,(x,y) in centers.items():
            if k in used:continue
            inside=(rect.x0-30<=x<=rect.x1+30 and rect.y0-4<=y<=rect.y1+4)
            if inside:
                dx=max(0,rect.x0-x,x-rect.x1);dy=max(0,rect.y0-y,y-rect.y1)
                possible.append((dx+dy,k))
        if possible:
            _,k=min(possible);out[k]=img;used.add(k)
    return out

def figure_assets(data):
    p=fitz.open(stream=data,filetype='pdf');assets={}
    for page in p:
        qs=_q_ranges(page);imgs=_small_images(page)
        for idx,(q,y0,y1) in enumerate(qs):
            next_y=qs[idx+1][1] if idx+1<len(qs) else page.rect.y1
            markers=_option_markers(page,y0,next_y)
            candidates=[]
            for r,x in imgs:
                if r.y1<=y1 or r.y0>=next_y or r.width>=page.rect.width*.95:continue
                candidates.append((r,_crop_image(page,r,pad=2,scale=2)))
            option_imgs={};figure_imgs=[]
            for r,img in candidates:
                split=_split_composite_option(page,r,markers)
                if split:
                    option_imgs.update({k:(r,v) for k,v in split.items()});continue
                assigned=_assign_option_images([(r,img)],markers)
                if assigned:
                    for k,v in assigned.items():option_imgs[k]=(r,v)
                else: figure_imgs.append((r,img))
            if option_imgs: assets.setdefault(q,{})['options']={k:v[1] for k,v in sorted(option_imgs.items()) if 1<=k<=4}
            if figure_imgs:
                r,img=max(figure_imgs,key=lambda z:z[0].width*z[0].height);assets.setdefault(q,{})['figure']=img
    return assets

@st.cache_data(show_spinner=False)
def assets_cached(data):return figure_assets(data)

@st.cache_data(show_spinner=False)
def ocr_image(img):
    try:
        from PIL import Image
        import pytesseract
        txt=pytesseract.image_to_string(Image.open(io.BytesIO(img)).convert('RGB'),config='--psm 6')
        return re.sub(r'\s+',' ',txt).strip()
    except Exception:return ''

def _api_key():
    key=os.getenv('OPENAI_API_KEY')
    if key:return key
    try:return st.secrets.get('OPENAI_API_KEY')
    except Exception:return None

@st.cache_data(show_spinner=False)
def vision_transcribe(img):
    key=_api_key()
    if not key:return ''
    try:
        b64=base64.b64encode(img).decode('ascii')
        payload={'model':'gpt-5.6-luna','input':[{'role':'user','content':[{'type':'input_text','text':'Transcribe this exam-paper visual into precise editable text for a question paper. Preserve every visible symbol, charge, subscript, superscript, bond, arrow, label, axis, scale and numerical value. If it is a chemical structure, describe the structure left-to-right with bond types and substituents; include a SMILES form when unambiguous, otherwise use [unclear] for unreadable parts. If it is a graph, state x/y axes, labels, intercepts, slope/curve shape, marked points and values. If it is a circuit, state every component and its connections. If it is a mathematical diagram, transcribe equations and labels exactly. Do not guess missing information. Return only the transcription, with no commentary.'},{'type':'input_image','image_url':f'data:image/png;base64,{b64}'}]}],'max_output_tokens':700}
        r=requests.post('https://api.openai.com/v1/responses',headers={'Authorization':f'Bearer {key}','Content-Type':'application/json'},json=payload,timeout=45);r.raise_for_status();j=r.json()
        if isinstance(j.get('output_text'),str):return j['output_text'].strip()
        parts=[]
        for item in j.get('output',[]):
            for c in item.get('content',[]):
                if isinstance(c,dict) and isinstance(c.get('text'),str):parts.append(c['text'])
        return re.sub(r'\s+',' ',' '.join(parts)).strip()
    except Exception:return ''

def safe_vision_note():return bool(_api_key())
def transcribe_visual(img):return vision_transcribe(img) or ocr_image(img)

@st.cache_data(show_spinner=False)
def _translate_google(text):
    for _ in range(2):
        try:
            y=GoogleTranslator(source='en',target='hi').translate(text)
            if y and y.strip(): return y.strip()
        except Exception:
            pass
    return ''

def _translate_mymemory(text):
    try:
        r=requests.get('https://api.mymemory.translated.net/get',params={'q':text,'langpair':'en|hi'},timeout=20)
        r.raise_for_status()
        y=((r.json().get('responseData') or {}).get('translatedText') or '').strip()
        if y and y.lower()!=text.strip().lower(): return y
    except Exception:
        pass
    return ''

def _restore_tokens(y,tokens):
    for i,v in enumerate(tokens):
        exact=f'PCX{i}X'
        patterns=(exact,f'[{exact}]',f'{{{exact}}}',f'<{exact}>',f'PCX {i} X',f'PCX{i} X',f'PCX {i}X')
        found=next((pat for pat in patterns if pat in y),None)
        if found is None:return None
        y=y.replace(found,v)
    return y

def _is_formula_only(text):
    if not text:return True
    return not re.search(r'[A-Za-z]{3,}',PROTECT.sub('',text))

def _translate_protected(text):
    tokens=[]
    def protect(m):
        tokens.append(m.group(0));return f'PCX{len(tokens)-1}X'
    masked=PROTECT.sub(protect,text)
    if not masked.strip():return text,True
    for method in (_translate_google,_translate_mymemory):
        y=method(masked)
        if not y:continue
        restored=_restore_tokens(y,tokens)
        if restored is None:continue
        if restored.strip()==text.strip() and not _is_formula_only(text) and len(re.findall(r'\b[A-Za-z]{2,}\b',text))>=4:continue
        return restored,True
    return '[Hindi translation unavailable]',False

@st.cache_data(show_spinner=False)
def tr(x):
    if not x:return x
    return _translate_protected(x)[0]

def _digits(text):return sorted(re.findall(r'\d+(?:\.\d+)?',text or ''))
def _symbol_counts(text):return {c:(text or '').count(c) for c in ('%','?','=','≤','≥','≠','→','←','±','∝')}
def _translation_issue(original,translated,label):
    if not original:return None
    if translated=='[Hindi translation unavailable]':return f'{label}: Hindi translation unavailable.'
    if not _is_formula_only(original) and translated.strip()==original.strip() and len(re.findall(r'\b[A-Za-z]{2,}\b',original))>=4:return f'{label}: English text appears unchanged.'
    if _digits(original)!=_digits(translated):return f'{label}: numerical value(s) changed or were lost.'
    a,b=_symbol_counts(original),_symbol_counts(translated)
    for k in a:
        if a[k]!=b[k]:return f'{label}: symbol/punctuation {k!r} changed.'
    return None

def translation_qa(rs):
    issues=[]
    for q in rs:
        items=[(f'Q{q["num"]}',q['stem'])]+[(f'Q{q["num"]} option {k}',o) for k,o in enumerate(q['options'],1) if o]
        for label,text in items:
            issue=_translation_issue(text,tr(text),label)
            if issue:issues.append(issue)
    return issues

def _add_typed_run(p,text,h=False,b=False):
    sub='₀₁₂₃₄₅₆₇₈₉';sup='⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁽⁾';pattern=re.compile(r'(\^[+-]?\d+|[₀₁₂₃₄₅₆₇₈₉]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁽⁾]+)')
    def style(r):
        r.font.name='Nirmala UI' if h else 'Calibri';r.font.size=Pt(9.5);r.bold=b;r._element.rPr.rFonts.set(qn('w:eastAsia'),'Nirmala UI' if h else 'Calibri')
    pos=0
    for m in pattern.finditer(text):
        if m.start()>pos:
            r=p.add_run(text[pos:m.start()]);style(r)
        token=m.group(0);r=p.add_run(token[1:] if token.startswith('^') else token);style(r)
        if token.startswith('^') or all(c in sup for c in token):r.font.superscript=True
        else:r.font.subscript=True
        pos=m.end()
    if pos<len(text):
        r=p.add_run(text[pos:]);style(r)

def run(p,t,h=False,b=False):
    p.paragraph_format.space_after=Pt(2);_add_typed_run(p,t,h,b)

def subject_map(data):
    p=fitz.open(stream=data,filetype='pdf');starts=[]
    for page in p:
        txt=page.get_text('text');found=None
        for name in ('PHYSICS','CHEMISTRY','BOTANY','ZOOLOGY','BIOLOGY','MATHEMATICS'):
            if re.search(rf'(?mi)^\s*{re.escape(name)}\s*$',txt):found=name.title()
        nums=[int(m.group(1)) for m in re.finditer(r'(?m)^\s*(\d{1,3})\s*\.',txt)]
        if found and nums:starts.append((nums[0],found))
    return sorted(starts)

def subject_for(q,starts):
    cur=None
    for n,s in starts:
        if q>=n:cur=s
        else:break
    return cur

def add_picture_paragraph(cell,img,width=2.7):
    p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(io.BytesIO(img),width=Inches(width))

def docx(rs,bar,assets=None,subjects=None):
    assets=assets or {};d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.35);s.left_margin=s.right_margin=Inches(.3)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;run(p,'BILINGUAL QUESTION PAPER',False,True)
    t=d.add_table(rows=1,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(('English','हिन्दी अनुवाद')):
        c=t.rows[0].cells[i];c.width=Inches(3.9);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;run(c.paragraphs[0],h,i==1,True)
    missing_visuals=[];last_subject=None
    for j,q in enumerate(rs):
        subj=subject_for(q['num'],subjects or [])
        if subj and subj!=last_subject:
            row=t.add_row();cell=row.cells[0].merge(row.cells[1]);run(cell.paragraphs[0],subj,False,True);last_subject=subj
        cs=t.add_row().cells;qa=assets.get(q['num'],{})
        for c in cs:c.width=Inches(3.9);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        for i,h in enumerate((False,True)):
            c=cs[i];qp=c.paragraphs[0];run(qp,f'{q["num"]}. ',h,True)
            if q['stem']:run(qp,tr(q['stem']) if h else q['stem'],h)
            if qa.get('figure'):add_picture_paragraph(c,qa['figure'])
            opts=[]
            for k,o in enumerate(q['options'],1):
                text=o
                if not text and qa.get('options',{}).get(k):
                    raw=transcribe_visual(qa['options'][k]);text=raw if raw else '[Image-based option: text/structure could not be reliably transcribed]'
                    if not raw:missing_visuals.append(q['num'])
                if text:opts.append((k,text))
            compact=bool(opts) and len(opts)>=3 and all(len(re.sub(r'\s+',' ',x).strip())<=45 for _,x in opts) and all(len(re.sub(r'\s+',' ',opts[i][1]).strip())+len(re.sub(r'\s+',' ',opts[i+1][1]).strip())<=82 for i in range(0,len(opts)-1,2))
            if compact:
                for z in range(0,len(opts),2):
                    p2=c.add_paragraph();p2.paragraph_format.space_after=Pt(0);k,text=opts[z];run(p2,f'({k}) '+(tr(text) if h else text),h)
                    if z+1<len(opts):k2,text2=opts[z+1];run(p2,'    ',h);run(p2,f'({k2}) '+(tr(text2) if h else text2),h)
            else:
                for k,text in opts:
                    p2=c.add_paragraph();p2.paragraph_format.space_after=Pt(0);run(p2,f'({k}) '+(tr(text) if h else text),h)
        bar.progress((j+1)/len(rs))
    b=io.BytesIO();d.save(b);return b.getvalue(),sorted(set(missing_visuals))

st.markdown('<div class="upload-shell"><div class="upload-title">Create Your Bilingual Question Paper</div><div class="upload-sub">Upload an English PDF and generate a clean, editable English + Hindi Word file.</div></div>',unsafe_allow_html=True)
up=st.file_uploader('Upload English Question Paper (PDF)',type='pdf')
if up:
    data=up.getvalue();rs=records(pages(data));assets=assets_cached(data);subjects=subject_map(data);a,b=st.columns(2);a.metric('Questions',len(rs));b.metric('Range',f'{rs[0]["num"]}–{rs[-1]["num"]}' if rs else '—')
    structure_issues=validate_records(rs)
    if not rs:st.error('No reliable numbered questions found. Scanned/image-only PDFs need OCR.')
    elif structure_issues:
        st.error('Question extraction needs review before Word generation:')
        for issue in structure_issues:st.write('• '+issue)
        st.warning('Generation is stopped to prevent a corrupted paper.')
    else:
        blank=sum(not o for q in rs for o in q['options']);visual_blank=sum(1 for q in rs for k in range(1,5) if not q['options'][k-1] and assets.get(q['num'],{}).get('options',{}).get(k))
        st.success(f'{len(rs)} questions detected. Sequence verified: 1–{rs[-1]["num"]}. Options will be typed as (1)–(4); no option images.')
        if blank:st.warning(f'{blank} option(s) have no extractable text. Image-based options are OCR-processed when possible; no option image is inserted.')
        if visual_blank and not safe_vision_note():st.info('Some image-based chemistry/graph options may still need a vision OCR service for exact structural transcription. The app will not guess or insert those option images.')
        if st.button('🚀 Generate Bilingual Word File',use_container_width=True):
            with st.spinner('Checking Hindi translations and scientific notation…'):
                translation_issues=translation_qa(rs)
            if translation_issues:
                st.error(f'Translation QA found {len(translation_issues)} issue(s). Word generation is stopped to prevent a corrupted bilingual paper.')
                for issue in translation_issues[:30]:st.write('• '+issue)
                if len(translation_issues)>30:st.write(f'• …and {len(translation_issues)-30} more issue(s).')
                st.session_state.pop('docx',None)
                st.session_state['translation_issues']=translation_issues
            else:
                st.success('Translation QA passed: protected notation, numbers and required symbols are intact.')
                st.session_state.pop('translation_issues',None)
                out,miss=docx(rs,st.progress(0),assets,subjects);st.session_state['docx']=out;st.session_state['missing_visuals']=miss
        if 'translation_issues' in st.session_state and st.session_state['translation_issues']:
            st.warning('Fix/ retry the translation before downloading. The previous Word output has been cleared.')
        if 'docx' in st.session_state:
            if st.session_state.get('missing_visuals'):st.warning('Unreliably readable image-based options were left as explicit placeholders instead of being hallucinated: '+', '.join(map(str,st.session_state['missing_visuals'])))
            st.download_button('📥 Download Word Document',st.session_state['docx'],'Bilingual_Question_Paper.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document',use_container_width=True)

st.markdown('<div class="footer">PaperCraft AI • English → English + Hindi • Editable Word Output</div>',unsafe_allow_html=True)
