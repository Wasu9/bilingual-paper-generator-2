"""Phase 10 regression tests for the final PaperCraft AI QA gate."""
from io import BytesIO

from docx import Document

from final_qa import final_docx_qa, final_paper_qa, question_structure_qa


def assert_ok(condition, message):
    if not condition:
        raise AssertionError(message)


def make_docx(question_number=1, include_all_options=True):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    en = table.cell(0, 0)
    hi = table.cell(0, 1)
    en.text = f"{question_number}. What is the value of x?"
    if include_all_options:
        en.add_paragraph("(1) 1 (2) 2")
        en.add_paragraph("(3) 3 (4) 4")
    else:
        en.add_paragraph("(1) 1 (2) 2")
    hi.text = "x का मान क्या है?"
    hi.add_paragraph("(1) 1 (2) 2")
    hi.add_paragraph("(3) 3 (4) 4")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def main():
    records = [{
        "num": 1,
        "stem": "What is x = 10?",
        "options": ["1", "2", "3", "4"],
    }]
    assert_ok(not question_structure_qa(records), "valid record should pass structural QA")

    errors, warnings = final_paper_qa(
        records,
        translate=lambda text: text.replace("What is", "क्या है").replace("x = 10", "x = 10"),
    )
    assert_ok(not errors, f"valid source record should pass final QA: {errors}")
    assert_ok(not warnings, f"valid source record should have no warnings: {warnings}")

    bad = [{"num": 1, "stem": "What is x?", "options": ["A", "B", "C"]}]
    errors, _ = final_paper_qa(bad)
    assert_ok(any("expected 4 option slots" in x for x in errors), "missing option structure was not detected")

    errors, warnings = final_docx_qa(make_docx())
    assert_ok(not errors, f"valid DOCX should pass final QA: {errors}")
    assert_ok(not warnings, f"valid DOCX should have no warnings: {warnings}")

    errors, _ = final_docx_qa(make_docx(include_all_options=False))
    assert_ok(any("does not contain exactly (1), (2), (3), (4)" in x for x in errors), "DOCX option loss was not detected")

    print("Phase 10 final QA suite: PASS")


if __name__ == "__main__":
    main()
