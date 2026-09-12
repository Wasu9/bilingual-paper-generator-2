"""Phase 11.2 regression tests for the real PDF -> parsed records -> DOCX path.

The test loads only the production DOCX-generation functions from app.py via AST,
so Streamlit is never started and no translation/network service is required.
"""
from __future__ import annotations

import ast
import io
import re
from pathlib import Path

from docx import Document


def _load_docx_namespace():
    app_path = Path(__file__).resolve().parents[1] / "app.py"
    tree = ast.parse(app_path.read_text(encoding="utf-8"))
    wanted = {"_add_typed_run", "run", "subject_for", "add_picture_paragraph", "docx"}
    nodes = []
    for node in tree.body:
        if isinstance(node, ast.Import) or isinstance(node, ast.ImportFrom):
            nodes.append(node)
        elif isinstance(node, ast.FunctionDef) and node.name in wanted:
            nodes.append(node)
    ns = {}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(app_path), "exec"), ns)

    # Deterministic test translation: preserve mathematical/numeric content while
    # making English prose visibly different from Hindi.  Numeric-only option text
    # is deliberately avoided because (1)-(4) are structural labels in the DOCX;
    # testing numeric integrity is covered by the question stem and Q4 formulas.
    def fake_tr(text):
        replacements = {
            "What is the value of x = 10?": "x = 10 का मान क्या है?",
            "Which quantity is conserved?": "कौन सी राशि संरक्षित रहती है?",
            "Mass": "द्रव्यमान",
            "Momentum": "संवेग",
            "Energy": "ऊर्जा",
            "Force": "बल",
            "E = mc²": "E = mc²",
            "E = mv²": "E = mv²",
            "E = m/c²": "E = m/c²",
            "F²t²/2m": "F²t²/2m",
            "F²t²/3m": "F²t²/3m",
            "Ft²/2m": "Ft²/2m",
            "Ft/2m": "Ft/2m",
        }
        return replacements.get(text, "हिन्दी: " + text)

    ns["tr"] = fake_tr
    return ns


class _Bar:
    def progress(self, value):
        return None


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


def _records():
    return [
        # Numeric integrity is tested in the stem, while options remain semantic
        # text so structural (1)-(4) labels cannot be confused with option values.
        {"num": 1, "stem": "What is the value of x = 10?", "options": ["Mass", "Momentum", "Energy", "Force"]},
        {"num": 2, "stem": "Which quantity is conserved?", "options": ["Mass", "Momentum", "Energy", "Force"]},
        # Long options must remain one-per-line.
        {"num": 3, "stem": "Choose the correct statement.", "options": [
            "This is a deliberately long option that should not be compacted into a shared line.",
            "Another deliberately long option that should remain on its own line in Word.",
            "A third long option used to verify stable vertical formatting in the output document.",
            "A fourth long option used to verify stable vertical formatting in the output document.",
        ]},
        # Superscript and mathematical notation must survive generation.
        {"num": 4, "stem": "Find E = mc².", "options": ["E = mc²", "E = mv²", "E = m/c²", "E = mc"]},
    ]


def main():
    ns = _load_docx_namespace()
    data, missing = ns["docx"](_records(), _Bar())
    _assert(isinstance(data, (bytes, bytearray)) and len(data) > 1000, "DOCX bytes were not generated")
    _assert(missing == [], f"unexpected visual transcription misses: {missing}")

    doc = Document(io.BytesIO(data))
    _assert(len(doc.tables) == 1, "generated DOCX should contain one bilingual table")
    table = doc.tables[0]
    _assert(len(table.columns) == 2, "bilingual table must have exactly two columns")
    _assert(table.rows[0].cells[0].text.strip() == "English", "English header missing")
    _assert("हिन्दी" in table.rows[0].cells[1].text, "Hindi header missing")

    question_rows = []
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        en = "\n".join(p.text for p in row.cells[0].paragraphs).strip()
        hi = "\n".join(p.text for p in row.cells[1].paragraphs).strip()
        if re.match(r"^\d+\.\s", en):
            question_rows.append((en, hi))

    _assert(len(question_rows) == 4, f"expected 4 generated question rows, found {len(question_rows)}")
    _assert([int(re.match(r"^(\d+)", en).group(1)) for en, _ in question_rows] == [1, 2, 3, 4], "question sequence/order changed in DOCX")

    for n, (en, hi) in enumerate(question_rows, 1):
        _assert(re.match(rf"^{n}\.\s+\S", en), f"Q{n}: question text is not on the same line as its number")
        _assert(sorted(set(re.findall(r"\((1|2|3|4)\)", en))) == ["1", "2", "3", "4"], f"Q{n}: option markers are not exactly (1)-(4)")
        _assert(not re.search(r"(?<![A-Za-z])\(?[A-Da-d]\)?[\).:]\s+", en), f"Q{n}: accidental A/B/C/D option marker")
        _assert(hi and hi != en, f"Q{n}: Hindi cell is empty or unchanged")

    # Verify Phase 5/7 save post-processing actually leaves deterministic OOXML
    # table geometry and non-splitting question rows when sitecustomize is loaded.
    import sitecustomize  # noqa: F401
    out = io.BytesIO()
    doc.save(out)
    check = Document(io.BytesIO(out.getvalue()))
    xml = check.tables[0]._tbl.xml
    _assert('w:type="fixed"' in xml, "table layout was not hardened to fixed")
    question_xml_rows = [r._tr.xml for r in check.tables[0].rows if re.search(r">\d+\.\s", r._tr.xml)]
    _assert(question_xml_rows, "question rows not found in OOXML")
    _assert(all("w:cantSplit" in x for x in question_xml_rows), "question rows are not protected from splitting")

    # Superscript run should be represented as a Word superscript property.
    superscript_found = False
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.superscript:
                superscript_found = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.superscript:
                            superscript_found = True
    _assert(superscript_found, "superscript notation was not encoded in the DOCX")

    # Import final QA only after generation, matching the production validation layer.
    from final_qa import final_docx_qa
    errors, warnings = final_docx_qa(data)
    _assert(not errors, f"generated DOCX failed final QA: {errors}")
    _assert(not warnings, f"generated DOCX produced unexpected warnings: {warnings}")

    print("Phase 11.2 DOCX end-to-end regression suite: PASS")


if __name__ == "__main__":
    main()
