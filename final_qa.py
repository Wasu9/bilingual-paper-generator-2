"""Final QA helpers for PaperCraft AI.

The module is intentionally side-effect free so it can be used both before
DOCX generation (source records) and after generation (the actual DOCX bytes).
"""
import io
import re

OPTION_RE = re.compile(r"\((1|2|3|4)\)")
LETTER_OPTION_RE = re.compile(r"(?<![A-Za-z])\(?([A-Da-d])\)?[\).:]\s+")
PLACEHOLDER_RE = re.compile(r"\[(?:Image-based option|unclear|Hindi translation unavailable)[^\]]*\]", re.I)


def _numbers(text):
    return sorted(re.findall(r"\d+(?:\.\d+)?", text or ""))


def _symbols(text):
    return {c: (text or "").count(c) for c in ("%", "=", "≤", "≥", "≠", "→", "←", "±", "∝")}


def question_structure_qa(records):
    """Return critical structural issues for parsed question records."""
    issues = []
    if not records:
        return ["No questions were generated."]

    nums = [int(q.get("num", 0)) for q in records]
    expected = list(range(1, max(nums) + 1))
    missing = [n for n in expected if n not in nums]
    duplicate = sorted({n for n in nums if nums.count(n) > 1})

    if nums[0] != 1:
        issues.append(f"Sequence starts at {nums[0]}, not 1.")
    if missing:
        issues.append("Missing question number(s): " + ", ".join(map(str, missing[:20])))
    if duplicate:
        issues.append("Duplicate question number(s): " + ", ".join(map(str, duplicate[:20])))
    if nums != sorted(nums):
        issues.append("Question numbers are out of order.")

    for q in records:
        n = q.get("num")
        stem = (q.get("stem") or "").strip()
        if not stem:
            issues.append(f"Q{n}: empty question text.")

        options = q.get("options") or []
        if len(options) != 4:
            issues.append(f"Q{n}: expected 4 option slots, found {len(options)}.")
        for k in range(min(4, len(options))):
            text = (options[k] or "").strip()
            if not text:
                issues.append(f"Q{n}: option ({k + 1}) is empty in extracted text.")
            if LETTER_OPTION_RE.search(text):
                issues.append(f"Q{n}: option ({k + 1}) contains an A/B/C/D-style marker.")

        whole = " ".join([stem] + [str(x or "") for x in options])
        if LETTER_OPTION_RE.search(whole):
            issues.append(f"Q{n}: accidental A/B/C/D option marker detected.")

    return issues


def translation_integrity_qa(records, translate):
    """Check Hindi output without requiring any particular translation service."""
    issues = []
    for q in records:
        n = q.get("num")
        items = [(f"Q{n}", q.get("stem") or "")]
        items += [(f"Q{n} option {i}", x or "") for i, x in enumerate(q.get("options") or [], 1) if x]
        for label, original in items:
            hindi = (translate(original) or "").strip()
            if not hindi:
                issues.append(f"{label}: Hindi translation is empty.")
                continue
            if hindi == "[Hindi translation unavailable]":
                issues.append(f"{label}: Hindi translation unavailable.")
                continue
            if re.search(r"[A-Za-z]{3,}", original) and hindi == original.strip():
                issues.append(f"{label}: English text appears unchanged.")
            if _numbers(original) != _numbers(hindi):
                issues.append(f"{label}: numerical value(s) changed or were lost.")
            a, b = _symbols(original), _symbols(hindi)
            for symbol, count in a.items():
                if count != b[symbol]:
                    issues.append(f"{label}: symbol {symbol!r} changed.")
                    break
    return issues


def visual_output_qa(records, assets=None, missing_visuals=None):
    """Report unresolved visual transcription problems as warnings."""
    assets = assets or {}
    missing = set(missing_visuals or [])
    warnings = []
    for q in records:
        n = q.get("num")
        qa = assets.get(n, {})
        for k, img in (qa.get("options") or {}).items():
            if not img:
                warnings.append(f"Q{n} option ({k}): visual asset is empty.")
        if n in missing:
            warnings.append(f"Q{n}: one or more image-based options could not be reliably transcribed.")
    return warnings


def placeholder_qa(records, assets=None):
    """Find unresolved explicit placeholders in generated source content."""
    issues = []
    for q in records:
        n = q.get("num")
        for label, text in [(f"Q{n}", q.get("stem") or "")] + [
            (f"Q{n} option {i}", x or "")
            for i, x in enumerate(q.get("options") or [], 1)
        ]:
            if PLACEHOLDER_RE.search(text):
                issues.append(f"{label}: unresolved placeholder present.")
    return issues


def final_paper_qa(records, translate=None, assets=None, missing_visuals=None):
    """Return (critical_errors, warnings) for the source-record pipeline."""
    errors = question_structure_qa(records)
    warnings = []
    if translate is not None:
        errors.extend(translation_integrity_qa(records, translate))
    warnings.extend(visual_output_qa(records, assets, missing_visuals))
    warnings.extend(placeholder_qa(records, assets))
    errors = list(dict.fromkeys(errors))
    warnings = [x for x in dict.fromkeys(warnings) if x not in errors]
    return errors, warnings


def _docx_text(cell):
    return "\n".join((p.text or "").strip() for p in cell.paragraphs if (p.text or "").strip()).strip()


def _content_without_question_number(text):
    """Remove the leading printed question number before cross-column QA."""
    return re.sub(r"^\s*\d{1,3}\.\s*", "", text or "", count=1)


def final_docx_qa(docx_bytes):
    """Validate the actual generated DOCX before it becomes downloadable.

    Returns ``(errors, warnings)``.  The checks intentionally operate on the
    rendered document structure rather than the pre-generation records, so this
    catches corruption introduced during Word generation/layout processing.
    """
    errors, warnings = [], []
    try:
        from docx import Document
        doc = Document(io.BytesIO(docx_bytes))
    except Exception as exc:
        return [f"Generated Word document could not be opened: {exc}"], []

    question_rows = []
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            en = _docx_text(row.cells[0])
            hi = _docx_text(row.cells[1])
            m = re.match(r"^\s*(\d{1,3})\.\s*", en)
            if not m:
                continue
            question_rows.append((int(m.group(1)), en, hi))

    if not question_rows:
        return ["Final DOCX QA: no question rows were found."], []

    nums = [n for n, _, _ in question_rows]
    expected = list(range(1, max(nums) + 1))
    missing = [n for n in expected if n not in nums]
    duplicate = sorted({n for n in nums if nums.count(n) > 1})
    if nums[0] != 1:
        errors.append(f"Final DOCX QA: sequence starts at {nums[0]}, not 1.")
    if missing:
        errors.append("Final DOCX QA: missing question number(s): " + ", ".join(map(str, missing[:20])))
    if duplicate:
        errors.append("Final DOCX QA: duplicate question number(s): " + ", ".join(map(str, duplicate[:20])))
    if nums != sorted(nums):
        errors.append("Final DOCX QA: question numbers are out of order.")

    for n, en, hi in question_rows:
        options = OPTION_RE.findall(en)
        unique_options = sorted(set(options), key=int)
        if unique_options != ["1", "2", "3", "4"]:
            errors.append(f"Q{n}: final DOCX does not contain exactly (1), (2), (3), (4).")
        if LETTER_OPTION_RE.search(en):
            errors.append(f"Q{n}: final DOCX contains an accidental A/B/C/D-style option marker.")
        if not en.strip():
            errors.append(f"Q{n}: final English question cell is empty.")
        if not hi.strip():
            errors.append(f"Q{n}: final Hindi question cell is empty.")
        elif re.search(r"[A-Za-z]{3,}", en) and hi.strip() == en.strip():
            errors.append(f"Q{n}: final Hindi question cell appears unchanged.")

        en_content = _content_without_question_number(en)
        if _numbers(en_content) != _numbers(hi):
            errors.append(f"Q{n}: final DOCX numerical value(s) changed or were lost between English and Hindi.")
        a, b = _symbols(en_content), _symbols(hi)
        for symbol, count in a.items():
            if count != b[symbol]:
                errors.append(f"Q{n}: final DOCX symbol {symbol!r} changed between English and Hindi.")
                break
        if PLACEHOLDER_RE.search(en) or PLACEHOLDER_RE.search(hi):
            warnings.append(f"Q{n}: final DOCX contains an explicit unresolved visual/translation placeholder.")

    return list(dict.fromkeys(errors)), list(dict.fromkeys(warnings))
